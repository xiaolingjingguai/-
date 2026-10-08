"""计算书导出：纯文本、Word（.docx，三线表）、Excel（.xlsx）。

计算书数据结构 report：
  title    标题，如“（二）液体泄漏速率计算”
  basis    依据条款，如“附录F.1.1 式(F.1)、表F.1”
  params   [[参数, 取值, 单位, 来源]]
  steps    [(计算项目, 公式及依据, 代入数值, 计算结果)]
  results  [结论句]
  warn     [提示]
  tables   可选，结果表 [(表名, 表头, 行, 列宽cm)]，置于“计算结果”中、结论段落之前
"""
STD = "《建设项目环境风险评价技术导则》（HJ 169-2018）"


def make_report(title, basis, params, steps, results, warn, tables=None):
    return {"title": title, "basis": basis, "params": [[str(x) for x in p] for p in params],
            "steps": [tuple(str(x) for x in s) for s in steps], "results": list(results), "warn": list(warn),
            "tables": [(c, list(h), [[str(x) for x in r] for r in rows], list(w)) for c, h, rows, w in (tables or [])]}


def to_text(rep):
    t = "%s\n依据：%s%s。\n1. 参数取值\n" % (rep["title"], STD, rep["basis"].strip())
    for i, p in enumerate(rep["params"], 1):
        t += "（%d）%s：%s%s（%s）\n" % (i, p[0], p[1], " " + p[2] if p[2] else "", p[3])
    t += "2. 计算过程\n"
    for i, s in enumerate(rep["steps"], 1):
        t += "（%d）%s：%s；代入 %s；得 %s\n" % (i, s[0], s[1], s[2], s[3])
    t += "3. 计算结果\n"
    for cap, head, rows, _ in rep.get("tables", []):
        t += cap + "\n" + "\t".join(head) + "\n" + "".join("\t".join(r) + "\n" for r in rows)
    t += "".join(r + "\n" for r in rep["results"])
    if rep["warn"]:
        t += "4. 提示\n" + "".join("（%d）%s\n" % (i, w) for i, w in enumerate(rep["warn"], 1))
    return t


# ---------------------------------------------------------------- Word
# 直接写 OOXML（标准库 zipfile），不依赖 python-docx／lxml，以减小打包体积。
from xml.sax.saxutils import escape as _esc

_CM = 567  # 1 cm = 567 twip


def _run(text, size=12, bold=False, east="宋体", west="Times New Roman", color=None):
    rpr = '<w:rFonts w:ascii="%s" w:hAnsi="%s" w:eastAsia="%s" w:cs="%s"/>' % (west, west, east, west)
    if bold:
        rpr += "<w:b/><w:bCs/>"
    if color:
        rpr += '<w:color w:val="%s"/>' % color
    rpr += '<w:sz w:val="%d"/><w:szCs w:val="%d"/>' % (round(size * 2), round(size * 2))
    return '<w:r><w:rPr>%s</w:rPr><w:t xml:space="preserve">%s</w:t></w:r>' % (rpr, _esc(text))


def _p(text, size=12, bold=False, east="宋体", indent=True, align=None, space_before=0, space_after=0, line=1.5, color=None):
    ppr = '<w:spacing w:before="%d" w:after="%d" w:line="%d" w:lineRule="auto"/>' % (space_before * 20, space_after * 20, round(line * 240))
    if indent:
        ppr += '<w:ind w:firstLine="%d"/>' % round(size * 2 * 20)
    if align:
        ppr += '<w:jc w:val="%s"/>' % align
    return "<w:p><w:pPr>%s</w:pPr>%s</w:p>" % (ppr, _run(text, size, bold, east, color=color) if text else "")


def _tbl(caption, header, rows, widths_cm):
    """三线表：顶线、底线 1.5 磅，表头下线 0.75 磅，无竖线。"""
    out = _p(caption, 10.5, True, "黑体", indent=False, align="center", space_before=6, line=1.0)
    tw = [round(w * _CM) for w in widths_cm]
    out += ('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/>'
            '<w:tblBorders><w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/><w:left w:val="nil"/>'
            '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/><w:right w:val="nil"/>'
            '<w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders><w:tblLayout w:type="fixed"/>'
            '<w:tblCellMar><w:left w:w="57" w:type="dxa"/><w:right w:w="57" w:type="dxa"/></w:tblCellMar></w:tblPr><w:tblGrid>%s</w:tblGrid>'
            % (sum(tw), "".join('<w:gridCol w:w="%d"/>' % w for w in tw)))
    for ri, vals in enumerate([header] + rows):
        out += "<w:tr>%s" % ("<w:trPr><w:tblHeader/></w:trPr>" if ri == 0 else "")
        for ci, v in enumerate(vals):
            tcpr = '<w:tcW w:w="%d" w:type="dxa"/>' % tw[ci]
            if ri == 0:
                tcpr += '<w:tcBorders><w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/></w:tcBorders>'
            tcpr += '<w:vAlign w:val="center"/>'
            ppr = '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
            if ri == 0 or ci == 0:
                ppr += '<w:jc w:val="center"/>'
            out += "<w:tc><w:tcPr>%s</w:tcPr><w:p><w:pPr>%s</w:pPr>%s</w:p></w:tc>" % (tcpr, ppr, _run(str(v), 10.5, ri == 0))
        out += "</w:tr>"
    return out + "</w:tbl>"


def _module_docx(rep, tno):
    b = _p(rep["title"], 14, True, "黑体", indent=False, space_before=12, space_after=6)
    b += _p("1. 计算依据", 12, True, indent=False)
    b += _p("依据%s%s。" % (STD, rep["basis"].strip()))
    b += _p("2. 参数取值", 12, True, indent=False)
    name = rep["title"].split("）", 1)[-1]
    if name.endswith("计算"):
        name = name[:-2]
    b += _tbl("表%d　%s参数取值" % (tno, name), ["序号", "参数", "取值", "单位", "取值依据／来源"],
              [[str(i), p[0], p[1], p[2] or "—", p[3]] for i, p in enumerate(rep["params"], 1)], [1.3, 3.8, 2.8, 1.6, 6.5])
    b += _p("3. 计算过程", 12, True, indent=False, space_before=6)
    b += _tbl("表%d　%s计算过程" % (tno + 1, name), ["序号", "计算项目", "公式及依据", "代入数值", "计算结果"],
              [[str(i), s[0], s[1], s[2], s[3]] for i, s in enumerate(rep["steps"], 1)], [1.3, 2.8, 4.9, 4.6, 2.4])
    b += _p("4. 计算结果", 12, True, indent=False, space_before=6)
    tno += 2
    for cap, head, rows, widths in rep.get("tables", []):
        b += _tbl("表%d　%s" % (tno, cap), head, rows, widths) + _p("", 6, indent=False, line=1.0)
        tno += 1
    for r in rep["results"]:
        b += _p(r)
    if rep["warn"]:
        b += _p("5. 编制提示（核对后删除）", 12, True, indent=False)
        for i, w in enumerate(rep["warn"], 1):
            b += _p("（%d）%s" % (i, w), 10.5, color="8A4B00")
    return b, tno


_CT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
       '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
       '<Default Extension="xml" ContentType="application/xml"/>'
       '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
       '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
       '</Types>')
_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
         '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
         '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
         '</Relationships>')
_DRELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
          '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
          '</Relationships>')
_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_STYLES = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles %s><w:docDefaults><w:rPrDefault><w:rPr>'
           '<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体" w:cs="Times New Roman"/>'
           '<w:sz w:val="24"/><w:szCs w:val="24"/><w:lang w:val="en-US" w:eastAsia="zh-CN"/></w:rPr></w:rPrDefault>'
           '<w:pPrDefault><w:pPr><w:spacing w:after="0"/></w:pPr></w:pPrDefault></w:docDefaults>'
           '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
           '<w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/><w:tblPr>'
           '<w:tblCellMar><w:left w:w="108" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr></w:style>'
           '</w:styles>') % _W


def to_docx(reports, path, doc_title="环境风险计算书"):
    import zipfile
    body = _p(doc_title, 16, True, "黑体", indent=False, align="center", space_after=6)
    body += _p("计算依据：%s；本计算书由 HJ 169 风险计算器生成，参数来源逐项列明，可复现。" % STD, 10.5, indent=False, align="center", space_after=6)
    tno = 1
    for rep in reports:
        b, tno = _module_docx(rep, tno)
        body += b
    sect = ('<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1417" w:right="1417" w:bottom="1417" w:left="1417"'
            ' w:header="851" w:footer="992" w:gutter="0"/></w:sectPr>')
    doc = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document %s><w:body>%s%s</w:body></w:document>' % (_W, body, sect)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _CT)
        z.writestr("_rels/.rels", _RELS)
        z.writestr("word/_rels/document.xml.rels", _DRELS)
        z.writestr("word/styles.xml", _STYLES)
        z.writestr("word/document.xml", doc)


# ---------------------------------------------------------------- Excel
def to_xlsx(reports, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hfill = PatternFill("solid", fgColor="E2F0EE")
    base = Font(name="宋体", size=10.5)
    bold = Font(name="宋体", size=10.5, bold=True)
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")
    summary = wb.create_sheet("汇总")
    summary.append(["模块", "计算结果"])
    for c in summary[1]:
        c.font, c.fill, c.border, c.alignment = bold, hfill, border, center
    for rep in reports:
        name = rep["title"].split("）", 1)[-1][:28]
        ws = wb.create_sheet(name)
        ws.column_dimensions["A"].width = 6
        for col, w in zip("BCDE", (22, 42, 36, 30)):
            ws.column_dimensions[col].width = w
        ws.append([rep["title"]])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        ws["A1"].font = Font(name="黑体", size=14, bold=True)
        ws.append(["依据：%s%s" % (STD, rep["basis"].strip())])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
        ws["A2"].font, ws["A2"].alignment = base, wrap
        ws.append([])

        def table(title, header, rows):
            ws.append([title])
            ws.cell(ws.max_row, 1).font = bold
            ws.append(header)
            for c in ws[ws.max_row]:
                c.font, c.fill, c.border, c.alignment = bold, hfill, border, center
            for r in rows:
                ws.append(r)
                for i, c in enumerate(ws[ws.max_row]):
                    c.font, c.border, c.alignment = base, border, center if i == 0 else wrap
            ws.append([])

        table("参数取值", ["序号", "参数", "取值", "单位", "取值依据／来源"],
              [[i, p[0], p[1], p[2] or "—", p[3]] for i, p in enumerate(rep["params"], 1)])
        table("计算过程", ["序号", "计算项目", "公式及依据", "代入数值", "计算结果"],
              [[i, s[0], s[1], s[2], s[3]] for i, s in enumerate(rep["steps"], 1)])
        for cap, head, rows, _ in rep.get("tables", []):
            table(cap, head, rows)
        ws.append(["计算结果"])
        ws.cell(ws.max_row, 1).font = bold
        for r in rep["results"]:
            ws.append([r])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=5)
            ws.cell(ws.max_row, 1).font, ws.cell(ws.max_row, 1).alignment = base, wrap
        if rep["warn"]:
            ws.append([])
            ws.append(["编制提示"])
            ws.cell(ws.max_row, 1).font = bold
            for w in rep["warn"]:
                ws.append([w])
                ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=5)
                ws.cell(ws.max_row, 1).font = Font(name="宋体", size=10.5, color="8A4B00")
                ws.cell(ws.max_row, 1).alignment = wrap
        summary.append([rep["title"], "；".join(rep["results"])])
        for i, c in enumerate(summary[summary.max_row]):
            c.font, c.border, c.alignment = base, border, wrap
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 100
    wb.save(path)
