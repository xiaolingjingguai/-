# -*- coding: utf-8 -*-
"""生态影响评价制图（桌面版）

依据《环境影响评价技术导则 生态影响》（HJ 19—2022）附录 D 及 7.4、8.2、9.1 等条款，
读入 shp／gpkg／gdb／dxf／kml／csv／xlsx 数据，批量生成生态影响评价图件并统计面积。
不依赖 ArcGIS。用法：
    python ecomap_gui.py                    打开界面
    python ecomap_gui.py --batch 配置.json   按配置直接出图
    python ecomap_gui.py --selftest 结果.txt 无人值守自检
"""
import copy
import json
import os

# 本机若装有 ArcGIS、QGIS 等，PROJ_LIB／GDAL_DATA 等环境变量可能指向不兼容版本，启动时清除，改用自带数据
for _k in ("PROJ_LIB", "PROJ_DATA", "GDAL_DATA", "GDAL_DRIVER_PATH"):
    os.environ.pop(_k, None)
import queue
import sys
import tempfile
import threading
import traceback
import tkinter as tk
from tkinter import colorchooser, filedialog, messagebox, ttk

import ecomap_core as C
import license_verify as LV

APP_NAME = "生态影响评价制图"
TITLE = "%s %s（依据 HJ 19—2022）" % (APP_NAME, C.VERSION)
VEC_TYPES = [("矢量数据", "*.shp *.gpkg *.geojson *.json *.kml *.kmz *.dxf *.gml *.tab *.mif"),
             ("点位表", "*.csv *.xlsx *.xls"), ("所有文件", "*.*")]
RAS_TYPES = [("栅格", "*.tif *.tiff *.img *.jp2 *.vrt"), ("所有文件", "*.*")]
EPSG_CHOICES = ["自动（按项目中心经度选 CGCS2000 3 度带）", "4548  CGCS2000 3度带 中央经线117°E",
                "4549  CGCS2000 3度带 中央经线120°E", "4547  CGCS2000 3度带 中央经线114°E",
                "4490  CGCS2000 经纬度（不推荐用于面积统计）"]
PT_EPSG = ["4490  CGCS2000 经纬度", "4326  WGS-84 经纬度", "4549  CGCS2000 3度带 120°E", "4548  CGCS2000 3度带 117°E"]

HELP = """一、软件用途
按《环境影响评价技术导则 生态影响》（HJ 19—2022）批量编制生态影响评价图件，并输出面积统计表与图件清单。
每幅图按附录 D.3 自动配置主图、图名、图例、比例尺（线段比例尺和数字比例尺）、方向标（指北针）、注记、
制图数据源、成图时间，另加公里网和图廓；比例尺按 D.2 取标准比例尺。

二、使用步骤
1. “项目信息”页：填写项目名称、评价等级、制图单位、图幅、输出目录；评价范围数据缺失时填写外扩距离（须按 6.2 论证）。
2. “图层数据”页：为各专题选择数据文件，并从下拉框选择分类字段、注记字段。
   · 支持 shp、gpkg、FileGDB（点“GDB”选文件夹）、geojson、kml、dxf；数据须带坐标系。
   · 点位也可用 csv/xlsx 表：点“表设置”指定经度/纬度（或 X/Y）列、坐标系；坐标若从高德、腾讯地图拾取（GCJ-02），勾选“GCJ-02 纠偏”。
3. “生态敏感区”页：每个敏感区添加一行，按 D.2 每处单独成图，并生成一幅总图（不同线型区分）。功能分区数据须用主管部门公布版本。
4. “底图与植被覆盖度”页：可选工作底图栅格（DOM/DRG GeoTIFF），或填写天地图 key 加参考底图；植被覆盖度可填 NDVI 栅格或红光、近红外波段。
5. “图件与输出”页：按评价等级显示各图件要求（●应编制 ○涉及时编制 △可选），勾选后“预览”或“生成全部”。
6. “配色”页：读取类别后可逐类改色，保存在配置文件中。

三、成果
输出目录下：各图件 PDF/JPG（300 dpi）、“生态图件清单及面积统计.xlsx”（图件清单、评价范围内与工程占用各类型面积及比例、植被覆盖度分级面积）、
“制图日志.txt”、“制图配置.ecomap.json”（可再次打开）。

四、需要复核的事项
1. 图件分级要求系按 HJ 19—2022 正文条款归纳；地理位置图、水系图归为通用图件属推断；2022 版已无 2011 版的分级图件对照表。
2. 工作底图应采用标准地形图（D.2），天地图等在线地图只宜作参考底图。
3. 植被覆盖度 NDVIs、NDVIv 取 5%/95% 分位与 30%/45%/60%/75% 分级断点为经验取值（分级参照 SL 190—2007），导则未规定，报告中须说明。
4. 土地利用（GB/T 21010 二级类）、生态系统（HJ 1166 Ⅱ级）分类由数据字段决定，本软件不做分类；默认配色为参考配色，非 TD/T 1055 等规定色值。
5. 面积按高斯-克吕格投影平面坐标计算。
6. 项目总平面布置图、施工布置图、线性工程平纵断面图、生态保护措施设计图取自设计文件，本软件不生成。
7. 本软件未经主管部门认定，成果须由编制人员复核后使用。"""


def parse_epsg(s):
    s = (s or "").strip()
    if not s or s.startswith("自动"):
        return ""
    return s.split()[0]


# ====================================================================== 通用控件
class ScrollFrame(ttk.Frame):
    def __init__(self, master, **kw):
        super().__init__(master, **kw)
        self.canvas = tk.Canvas(self, highlightthickness=0, borderwidth=0)
        self.vsb = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.hsb = ttk.Scrollbar(self, orient="horizontal", command=self.canvas.xview)
        self.inner = ttk.Frame(self.canvas, padding=(6, 6, 12, 6))
        self.inner.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.canvas.configure(yscrollcommand=self.vsb.set, xscrollcommand=self.hsb.set)
        self.vsb.pack(side="right", fill="y")
        self.hsb.pack(side="bottom", fill="x")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.inner.bind("<Enter>", lambda e: self.canvas.bind_all("<MouseWheel>", self._wheel))
        self.inner.bind("<Leave>", lambda e: self.canvas.unbind_all("<MouseWheel>"))

    def _wheel(self, e):
        self.canvas.yview_scroll(int(-e.delta / 120), "units")


class LayerRow:
    """一个数据槽：路径、图层、字段、注记、海拔、数据源、点位表设置。"""

    def __init__(self, app, parent, row, key, name, needs, removable=False):
        self.app, self.key, self.name = app, key, name
        self.v = {k: tk.StringVar() for k in ("path", "layer", "field", "label", "elev", "source", "x", "y", "epsg",
                                              "name")}
        self.gcj = tk.BooleanVar(value=False)
        self.widgets = []
        c = 0
        if key == "sensitive":
            e = ttk.Entry(parent, textvariable=self.v["name"], width=16)
            e.grid(row=row, column=c, sticky="we", padx=2, pady=2)
        else:
            e = ttk.Label(parent, text=name)
            e.grid(row=row, column=c, sticky="w", padx=2, pady=2)
        self.widgets.append(e)
        c += 1
        e = ttk.Entry(parent, textvariable=self.v["path"], width=40)
        e.grid(row=row, column=c, sticky="we", padx=2)
        e.bind("<FocusOut>", lambda ev: self.refresh_fields())
        self.widgets.append(e)
        c += 1
        b = ttk.Button(parent, text="文件", width=5, command=self.browse)
        b.grid(row=row, column=c, padx=1)
        self.widgets.append(b)
        c += 1
        b = ttk.Button(parent, text="GDB", width=4, command=self.browse_gdb)
        b.grid(row=row, column=c, padx=1)
        self.widgets.append(b)
        c += 1
        self.cb_layer = ttk.Combobox(parent, textvariable=self.v["layer"], width=12)
        self.cb_layer.grid(row=row, column=c, padx=2)
        self.cb_layer.bind("<<ComboboxSelected>>", lambda ev: self.refresh_fields(keep_layer=True))
        self.widgets.append(self.cb_layer)
        c += 1
        self.cbs = {}
        for k in ("field", "label", "elev"):
            cb = ttk.Combobox(parent, textvariable=self.v[k], width=11)
            if k in needs or (key == "sensitive" and k == "field"):
                cb.grid(row=row, column=c, padx=2)
                self.widgets.append(cb)
            self.cbs[k] = cb
            c += 1
        e = ttk.Entry(parent, textvariable=self.v["source"], width=30)
        e.grid(row=row, column=c, sticky="we", padx=2)
        self.widgets.append(e)
        c += 1
        b = ttk.Button(parent, text="表设置", width=6, command=self.table_dialog)
        b.grid(row=row, column=c, padx=1)
        self.widgets.append(b)
        c += 1
        if removable:
            b = ttk.Button(parent, text="删除", width=5, command=lambda: app.remove_sensitive(self))
            b.grid(row=row, column=c, padx=1)
            self.widgets.append(b)
        else:
            b = ttk.Button(parent, text="清空", width=5, command=self.clear)
            b.grid(row=row, column=c, padx=1)
            self.widgets.append(b)

    def destroy(self):
        for w in self.widgets:
            w.destroy()

    def clear(self):
        for k in ("path", "layer", "field", "label", "elev", "x", "y"):
            self.v[k].set("")
        self.gcj.set(False)

    def browse(self):
        p = filedialog.askopenfilename(title="选择%s数据" % self.name, filetypes=VEC_TYPES)
        if p:
            self.v["path"].set(os.path.normpath(p))
            self.v["layer"].set("")
            self.refresh_fields()
            if p.lower().endswith((".csv", ".xlsx", ".xls")):
                self.table_dialog()

    def browse_gdb(self):
        p = filedialog.askdirectory(title="选择 FileGDB（.gdb 文件夹）")
        if p:
            self.v["path"].set(os.path.normpath(p))
            self.v["layer"].set("")
            self.refresh_fields()

    def refresh_fields(self, keep_layer=False):
        p = self.v["path"].get().strip()
        if not p or not os.path.exists(p):
            return
        try:
            if not p.lower().endswith((".csv", ".xlsx", ".xls", ".txt")):
                lays = C.list_layers(p)
                self.cb_layer["values"] = lays
                if lays and (not keep_layer and not self.v["layer"].get() or self.v["layer"].get() not in lays):
                    self.v["layer"].set(lays[0])
            fields = [""] + C.list_fields(p, self.v["layer"].get())
            for cb in self.cbs.values():
                cb["values"] = fields
            self.fields = fields
        except Exception as e:
            self.app.log("读取字段失败：%s（%s）" % (p, e))

    def table_dialog(self):
        p = self.v["path"].get().strip()
        if not p.lower().endswith((".csv", ".xlsx", ".xls", ".txt")):
            messagebox.showinfo(APP_NAME, "“表设置”仅用于 csv/xlsx 点位表。矢量数据自带坐标系，无需设置。")
            return
        self.refresh_fields()
        d = tk.Toplevel(self.app.root)
        d.title("点位表设置：%s" % self.name)
        d.transient(self.app.root)
        f = ttk.Frame(d, padding=12)
        f.pack(fill="both", expand=True)
        cols = getattr(self, "fields", [""])
        ttk.Label(f, text="X 列（经度或东坐标）").grid(row=0, column=0, sticky="w", pady=3)
        ttk.Combobox(f, textvariable=self.v["x"], values=cols, width=18).grid(row=0, column=1, pady=3)
        ttk.Label(f, text="Y 列（纬度或北坐标）").grid(row=1, column=0, sticky="w", pady=3)
        ttk.Combobox(f, textvariable=self.v["y"], values=cols, width=18).grid(row=1, column=1, pady=3)
        ttk.Label(f, text="坐标系").grid(row=2, column=0, sticky="w", pady=3)
        ep = tk.StringVar(value=next((s for s in PT_EPSG if s.startswith(self.v["epsg"].get() or "4490")),
                                     self.v["epsg"].get()))
        ttk.Combobox(f, textvariable=ep, values=PT_EPSG, width=30).grid(row=2, column=1, pady=3)
        ttk.Checkbutton(f, text="坐标从高德/腾讯地图拾取（GCJ-02），纠偏到 WGS-84", variable=self.gcj).grid(
            row=3, column=0, columnspan=2, sticky="w", pady=3)
        ttk.Label(f, text="说明：3 度带平面坐标的 X 填东坐标（不带带号，约 5 位整数加 50 万），Y 填北坐标。",
                  foreground="#666").grid(row=4, column=0, columnspan=2, sticky="w")

        def ok():
            self.v["epsg"].set(ep.get().split()[0] if ep.get().strip() else "4490")
            d.destroy()
        ttk.Button(f, text="确定", command=ok).grid(row=5, column=1, sticky="e", pady=8)

    def get(self):
        d = {k: self.v[k].get().strip() for k in ("path", "layer", "field", "label", "elev", "source", "x", "y",
                                                   "epsg")}
        d["gcj02"] = bool(self.gcj.get())
        if self.key == "sensitive":
            d["name"] = self.v["name"].get().strip()
        return d

    def set(self, d):
        for k in self.v:
            if k in d and d[k] is not None:
                self.v[k].set(str(d[k]))
        self.gcj.set(bool(d.get("gcj02")))
        self.app.root.after(10, lambda: self.refresh_fields(keep_layer=True))


# ====================================================================== 主程序
class App:
    def __init__(self, root):
        self.root = root
        root.title(TITLE)
        root.geometry("1280x800")
        root.minsize(980, 640)
        self.q = queue.Queue()
        self.busy = False
        self.cfg = C.default_config()
        self.cfg_path = ""
        self.ctx_cache = (None, None)
        style = ttk.Style()
        try:
            style.theme_use("vista" if sys.platform == "win32" else "clam")
        except tk.TclError:
            pass
        style.configure("Treeview", rowheight=24)
        style.configure("Head.TLabel", font=("Microsoft YaHei", 10, "bold"))

        top = ttk.Frame(root, padding=(8, 6))
        top.pack(fill="x")
        for t, cmd in (("新建", self.new), ("打开配置…", self.open_cfg), ("保存配置", self.save_cfg),
                       ("另存为…", self.save_cfg_as), ("载入示例项目", self.load_sample)):
            ttk.Button(top, text=t, command=cmd).pack(side="left", padx=3)
        self.lbl_cfg = ttk.Label(top, text="未保存的配置", foreground="#666")
        self.lbl_cfg.pack(side="left", padx=12)
        ttk.Button(top, text="生成全部图件", command=self.run_all).pack(side="right", padx=3)

        self.nb = ttk.Notebook(root)
        self.nb.pack(fill="both", expand=True, padx=8, pady=(0, 4))
        self._tab_info()
        self._tab_layers()
        self._tab_sensitive()
        self._tab_raster()
        self._tab_maps()
        self._tab_colors()
        self._tab_data()
        self._tab_license()
        self._tab_help()

        bot = ttk.Frame(root, padding=(8, 2, 8, 6))
        bot.pack(fill="x")
        self.pb = ttk.Progressbar(bot, mode="determinate", length=260)
        self.pb.pack(side="left")
        self.lbl_status = ttk.Label(bot, text="就绪")
        self.lbl_status.pack(side="left", padx=10)
        self.set_cfg(self.cfg)
        root.after(100, self._poll)

    # ------------------------------------------------------------------ 页面
    def _tab_info(self):
        f = ttk.Frame(self.nb, padding=14)
        self.nb.add(f, text="项目信息")
        self.iv = {k: tk.StringVar() for k in ("project_name", "eval_level", "mapper", "page", "dpi", "epsg",
                                               "buffer_m", "height_datum", "out_dir", "scale_mode")}
        self.fmt = {k: tk.BooleanVar() for k in ("pdf", "jpg", "png")}
        rows = [("项目名称", "project_name", "entry", None),
                ("生态影响评价等级", "eval_level", "combo", C.LEVELS),
                ("制图单位", "mapper", "entry", None),
                ("图幅", "page", "combo", list(C.PAGES)),
                ("比例尺取值", "scale_mode", "combo", ["标准比例尺", "充满图幅"]),
                ("JPG/PNG 分辨率（dpi）", "dpi", "combo", ["150", "200", "300", "400", "600"]),
                ("平面坐标系", "epsg", "combo", EPSG_CHOICES),
                ("评价范围外扩距离（m）", "buffer_m", "entry", None),
                ("高程基准", "height_datum", "entry", None),
                ("输出目录", "out_dir", "dir", None)]
        for i, (lab, k, kind, vals) in enumerate(rows):
            ttk.Label(f, text=lab).grid(row=i, column=0, sticky="w", pady=5, padx=(0, 10))
            if kind == "combo":
                w = ttk.Combobox(f, textvariable=self.iv[k], values=vals, width=50,
                                 state="readonly" if k in ("eval_level", "page", "scale_mode") else "normal")
                w.grid(row=i, column=1, sticky="w")
                if k == "eval_level":
                    w.bind("<<ComboboxSelected>>", lambda e: self.level_changed())
            else:
                ttk.Entry(f, textvariable=self.iv[k], width=60).grid(row=i, column=1, sticky="w")
            if kind == "dir":
                ttk.Button(f, text="浏览", command=self.pick_out).grid(row=i, column=2, padx=4)
        i = len(rows)
        ttk.Label(f, text="输出格式").grid(row=i, column=0, sticky="w", pady=5)
        ff = ttk.Frame(f)
        ff.grid(row=i, column=1, sticky="w")
        for k in ("pdf", "jpg", "png"):
            ttk.Checkbutton(ff, text=k.upper(), variable=self.fmt[k]).pack(side="left", padx=(0, 12))
        notes = ("说明：\n"
                 "1. 评价等级按 HJ 19—2022 6.1 判定，本软件只据此列出应编制图件，不做等级判定。\n"
                 "2. “图层数据”页提供评价范围数据时，外扩距离不起作用；外扩距离须按 6.2 论证，导则未给一般项目固定距离。\n"
                 "3. 平面坐标系选“自动”时，按项目中心经度选 CGCS2000 3 度带：泉州东部多为中央经线 120°E（EPSG 4549），\n"
                 "   安溪、永春、德化等西部多为 117°E（EPSG 4548），请以项目所在位置核对。\n"
                 "4. 比例尺“标准比例尺”取 1:2000、1:5000、1:10000、1:25000、1:50000 等（D.2）；“充满图幅”取整到 1、1.2、1.5、2、2.5…×10ⁿ。")
        ttk.Label(f, text=notes, foreground="#555", justify="left").grid(row=i + 1, column=0, columnspan=3,
                                                                          sticky="w", pady=(16, 0))

    def _tab_layers(self):
        outer = ttk.Frame(self.nb)
        self.nb.add(outer, text="图层数据")
        sf = ScrollFrame(outer)
        sf.pack(fill="both", expand=True)
        p = sf.inner
        heads = ["专题", "数据路径", "", "", "图层", "分类字段", "注记字段", "海拔字段", "制图数据源（写入图面注记）", "", ""]
        for c, h in enumerate(heads):
            ttk.Label(p, text=h, style="Head.TLabel").grid(row=0, column=c, sticky="w", padx=2, pady=(0, 4))
        self.rows = {}
        for i, (key, name, needs, _) in enumerate(C.LAYER_SLOTS, 1):
            self.rows[key] = LayerRow(self, p, i, key, name, needs)
        ttk.Label(p, text="提示：shp 的字段名、属性为中文时请确保有 .cpg 文件（GBK 编码的 shp 若乱码，可在 GIS 软件中另存为 gpkg）。"
                          "项目占地范围可含永久、临时占地多个面。", foreground="#666").grid(
            row=len(C.LAYER_SLOTS) + 2, column=0, columnspan=10, sticky="w", pady=(10, 0))

    def _tab_sensitive(self):
        outer = ttk.Frame(self.nb, padding=(0, 6))
        self.nb.add(outer, text="生态敏感区")
        bar = ttk.Frame(outer, padding=(8, 0))
        bar.pack(fill="x")
        ttk.Button(bar, text="添加敏感区", command=self.add_sensitive).pack(side="left")
        ttk.Label(bar, text="  每处敏感区单独成图（D.2），并生成总图；功能分区字段可空。数据须用主管部门公布的功能区划。",
                  foreground="#555").pack(side="left")
        sf = ScrollFrame(outer)
        sf.pack(fill="both", expand=True)
        self.sens_parent = sf.inner
        heads = ["名称", "数据路径", "", "", "图层", "功能分区字段", "", "", "制图数据源", "", ""]
        for c, h in enumerate(heads):
            ttk.Label(self.sens_parent, text=h, style="Head.TLabel").grid(row=0, column=c, sticky="w", padx=2)
        self.sens_rows = []
        self.sens_next = 1

    def add_sensitive(self, d=None):
        r = LayerRow(self, self.sens_parent, self.sens_next, "sensitive", "生态敏感区", [], removable=True)
        self.sens_next += 1
        if d:
            r.set(d)
        else:
            r.v["name"].set("生态敏感区%d" % (len(self.sens_rows) + 1))
            r.v["source"].set("主管部门公布的功能区划图")
        self.sens_rows.append(r)

    def remove_sensitive(self, r):
        r.destroy()
        self.sens_rows.remove(r)

    def _tab_raster(self):
        f = ttk.Frame(self.nb, padding=14)
        self.nb.add(f, text="底图与植被覆盖度")
        self.bv = {k: tk.StringVar() for k in ("path", "source", "tianditu_key", "tianditu_type")}
        self.fv = {k: tk.StringVar() for k in ("ndvi", "red", "nir", "scale", "offset", "s_pct", "v_pct", "breaks",
                                               "labels", "source")}
        ttk.Label(f, text="一、工作底图（可选）", style="Head.TLabel").grid(row=0, column=0, sticky="w", pady=(0, 6))
        r = 1
        for lab, k, browse in (("底图栅格（DOM/DRG，GeoTIFF）", "path", True), ("底图数据源", "source", False)):
            ttk.Label(f, text=lab).grid(row=r, column=0, sticky="w", pady=3)
            ttk.Entry(f, textvariable=self.bv[k], width=70).grid(row=r, column=1, sticky="w")
            if browse:
                ttk.Button(f, text="浏览", command=lambda k=k: self._pick(self.bv[k], RAS_TYPES)).grid(row=r, column=2)
            r += 1
        ttk.Label(f, text="或：天地图参考底图").grid(row=r, column=0, sticky="w", pady=3)
        ff = ttk.Frame(f)
        ff.grid(row=r, column=1, sticky="w")
        ttk.Combobox(ff, textvariable=self.bv["tianditu_type"], values=["不使用"] + list(C.TIANDITU_TYPES),
                     state="readonly", width=16).pack(side="left")
        ttk.Label(ff, text="  key：").pack(side="left")
        ttk.Entry(ff, textvariable=self.bv["tianditu_key"], width=36, show="*").pack(side="left")
        r += 1
        ttk.Label(f, text="说明：D.2 要求以标准地形图为工作底图；天地图需自行在天地图网站申请 key、需联网，只宜作参考底图。"
                          "底图只加在地理位置图、水系图、点位类专题图和敏感区图上，土地利用等面状专题图不加。",
                  foreground="#555", wraplength=900, justify="left").grid(row=r, column=0, columnspan=3, sticky="w")
        r += 1
        ttk.Separator(f).grid(row=r, column=0, columnspan=3, sticky="we", pady=12)
        r += 1
        ttk.Label(f, text="二、植被覆盖度（附录 C.8.1，式 C.5）", style="Head.TLabel").grid(row=r, column=0, sticky="w")
        r += 1
        items = [("NDVI 栅格（与下方二选一）", "ndvi", True), ("红光波段（Sentinel-2 B4／Landsat 8/9 B4）", "red", True),
                 ("近红外波段（Sentinel-2 B8／Landsat 8/9 B5）", "nir", True),
                 ("反射率换算系数 scale", "scale", False), ("反射率换算偏移 offset", "offset", False),
                 ("NDVIs 取值百分位（%）", "s_pct", False), ("NDVIv 取值百分位（%）", "v_pct", False),
                 ("分级断点（0～1，逗号分隔）", "breaks", False), ("分级名称（逗号分隔，比断点多 1 个）", "labels", False),
                 ("影像数据源（写入图面注记）", "source", False)]
        for lab, k, browse in items:
            ttk.Label(f, text=lab).grid(row=r, column=0, sticky="w", pady=3)
            ttk.Entry(f, textvariable=self.fv[k], width=70).grid(row=r, column=1, sticky="w")
            if browse:
                ttk.Button(f, text="浏览", command=lambda k=k: self._pick(self.fv[k], RAS_TYPES)).grid(row=r, column=2)
            r += 1
        ttk.Label(f, text="说明：ρ = DN × scale + offset。Landsat Collection 2 L2 为 0.0000275、-0.2；"
                          "Sentinel-2 L2A（处理基线 04.00 起）为 0.0001、-0.1，请按影像元数据核对。\n"
                          "NDVIs、NDVIv 百分位与分级断点导则未规定，默认值为经验取值（分级参照 SL 190—2007），报告中须说明取值依据。",
                  foreground="#555", justify="left").grid(row=r, column=0, columnspan=3, sticky="w", pady=(6, 0))

    def _tab_maps(self):
        f = ttk.Frame(self.nb, padding=8)
        self.nb.add(f, text="图件与输出")
        pw = ttk.PanedWindow(f, orient="horizontal")
        pw.pack(fill="both", expand=True)
        left = ttk.Frame(pw)
        pw.add(left, weight=3)
        cols = ("sel", "title", "req", "basis", "data")
        tv = ttk.Treeview(left, columns=cols, show="headings", height=18)
        for c, t, w in zip(cols, ("生成", "图件", "要求", "HJ 19—2022 条款", "数据"), (50, 230, 80, 280, 200)):
            tv.heading(c, text=t)
            tv.column(c, width=w, anchor="center" if c in ("sel", "req") else "w")
        tv.pack(fill="both", expand=True)
        tv.bind("<Button-1>", self._toggle_map)
        tv.bind("<Double-1>", lambda e: self.preview())
        self.tv = tv
        self.lbl_req = ttk.Label(left, text="", foreground="#555", justify="left", wraplength=860)
        self.lbl_req.pack(fill="x", pady=4)
        bar = ttk.Frame(left)
        bar.pack(fill="x")
        for t, cmd in (("只选必须图件", lambda: self._select_maps("level")), ("全选", lambda: self._select_maps("all")),
                       ("全不选", lambda: self._select_maps("none")),
                       ("预览所选行", self.preview), ("生成全部勾选图件", self.run_all),
                       ("打开输出目录", self.open_out)):
            ttk.Button(bar, text=t, command=cmd).pack(side="left", padx=3)
        right = ttk.LabelFrame(pw, text="运行日志", padding=4)
        pw.add(right, weight=2)
        self.txt = tk.Text(right, wrap="word", width=40, font=("Microsoft YaHei", 9))
        sb = ttk.Scrollbar(right, command=self.txt.yview)
        self.txt.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.txt.pack(fill="both", expand=True)
        self.map_sel = {s["id"]: True for s in C.MAP_SPECS}

    def _tab_colors(self):
        f = ttk.Frame(self.nb, padding=10)
        self.nb.add(f, text="配色")
        bar = ttk.Frame(f)
        bar.pack(fill="x")
        ttk.Label(bar, text="专题：").pack(side="left")
        self.color_key = tk.StringVar(value="landuse")
        keys = ["landuse", "vegetation", "ecosystem", "targets", "habitat", "migration", "samples", "measures", "monitor"]
        self.color_keys = {C.SLOT_NAMES[k]: k for k in keys}
        self.color_name = tk.StringVar(value=C.SLOT_NAMES["landuse"])
        cb = ttk.Combobox(bar, textvariable=self.color_name, values=list(self.color_keys), state="readonly", width=20)
        cb.pack(side="left")
        ttk.Button(bar, text="读取类别", command=self.load_categories).pack(side="left", padx=6)
        ttk.Button(bar, text="恢复参考配色", command=self.reset_colors).pack(side="left", padx=6)
        ttk.Label(bar, text="  双击类别改色。默认为参考配色（按类别名称关键词），非标准规定色值。", foreground="#555").pack(side="left")
        self.ctv = ttk.Treeview(f, columns=("cat", "color"), show="tree headings", height=20)
        self.ctv.heading("#0", text="色块")
        self.ctv.column("#0", width=70)
        self.ctv.heading("cat", text="类别")
        self.ctv.column("cat", width=320)
        self.ctv.heading("color", text="颜色")
        self.ctv.column("color", width=120)
        self.ctv.pack(fill="both", expand=True, pady=6)
        self.ctv.bind("<Double-1>", self.edit_color)
        self._swatches = {}

    def _tab_data(self):
        import ecomap_links as K
        f = ttk.Frame(self.nb, padding=10)
        self.nb.add(f, text="数据获取")
        ttk.Label(f, text="一、可在线获取的数据（双击打开网站）", style="Head.TLabel").pack(anchor="w")
        cols = ("cat", "site", "org", "data", "use", "url")
        tv = ttk.Treeview(f, columns=cols, show="headings", height=len(K.LINKS))
        for c, t, w in zip(cols, ("类别", "网站", "主办单位", "可获取的数据", "用于", "网址"),
                           (70, 200, 170, 330, 140, 230)):
            tv.heading(c, text=t)
            tv.column(c, width=w, anchor="w")
        for i, (cat, data, site, org, url, use) in enumerate(K.LINKS):
            tv.insert("", "end", iid=str(i), values=(cat, site, org, data, use, url))
        tv.bind("<Double-1>", lambda e: self._open_url(tv, e))
        tv.pack(fill="x", pady=(4, 2))
        ttk.Label(f, text="说明：以上网址于 %s 访问核实可打开，主办单位已核对；是否需要注册登录、数据范围以网站当前情况为准。"
                          "“科研机构”类不是政府官方发布的数据，只宜作参考或辅助解译。\n"
                          "以下站点未能核实网址，故未收录（信息不足）：%s。" % (K.VERIFIED_DATE, "、".join(K.NOT_VERIFIED)),
                  foreground="#555", wraplength=1100, justify="left").pack(anchor="w")
        ttk.Label(f, text="二、不公开下载、须向主管部门申请的数据", style="Head.TLabel").pack(anchor="w", pady=(12, 0))
        tv2 = ttk.Treeview(f, columns=("data", "where", "use"), show="headings", height=len(K.APPLY))
        for c, t, w in zip(("data", "where", "use"), ("数据", "获取途径", "用于"), (360, 520, 200)):
            tv2.heading(c, text=t)
            tv2.column(c, width=w, anchor="w")
        for r in K.APPLY:
            tv2.insert("", "end", values=r)
        tv2.pack(fill="x", pady=(4, 2))
        ttk.Label(f, text="说明：以上数据通常无公开下载网址（信息不足以提供链接），须按程序向主管部门申请；"
                          "1:10000、1:50000 地形图等测绘成果按国家涉密测绘成果管理规定使用，D.2 要求以标准地形图作工作底图。",
                  foreground="#555", wraplength=1100, justify="left").pack(anchor="w")

    def _open_url(self, tv, e):
        import webbrowser
        iid = tv.identify_row(e.y)
        if iid:
            webbrowser.open(tv.set(iid, "url"))

    def _tab_license(self):
        f = ttk.Frame(self.nb, padding=14)
        self.nb.add(f, text="软件授权")
        ttk.Label(f, text="软件授权（机器绑定，离线使用）", style="Head.TLabel").grid(
            row=0, column=0, columnspan=3, sticky="w", pady=(0, 6))
        state, payload, reason = LV.license_state()
        if state == "licensed":
            name = payload.get("name") or "（未填名称）"
            exp = payload.get("expiry") or "永久"
            stxt = "已授权：授权给 %s，到期 %s" % (name, exp)
            scol = "#1a7f37"
        elif state == "trial":
            mins = max(1, (int(payload) + 59) // 60)
            stxt = "未授权，试用中：本次剩余约 %d 分钟" % mins
            scol = "#9a6a00"
        else:
            stxt = "未授权，试用已用完：%s" % reason
            scol = "#b00020"
        ttk.Label(f, text="当前状态：" + stxt, foreground=scol).grid(row=1, column=0, columnspan=3, sticky="w",
                                                                   pady=(0, 10))
        ttk.Label(f, text="本机机器码（请复制后发给软件作者申请授权）：").grid(row=2, column=0, columnspan=3, sticky="w")
        self._lic_code = tk.StringVar(value=LV.machine_code())
        ent = ttk.Entry(f, textvariable=self._lic_code, width=26, font=("Consolas", 13), justify="center")
        ent.grid(row=3, column=0, sticky="w", pady=(4, 6))
        ent.configure(state="readonly")

        def copy_code():
            try:
                self.root.clipboard_clear()
                self.root.clipboard_append(self._lic_code.get())
                bcopy.configure(text="已复制 ✓")
                self.root.after(1500, lambda: bcopy.configure(text="复制机器码"))
            except Exception:
                pass

        bcopy = ttk.Button(f, text="复制机器码", command=copy_code)
        bcopy.grid(row=3, column=1, sticky="w", padx=6)

        def import_license():
            p = filedialog.askopenfilename(title="选择授权文件 license.key",
                                           filetypes=[("授权文件", "*.key"), ("全部", "*.*")])
            if not p:
                return
            try:
                with open(p, encoding="utf-8") as fh:
                    raw = fh.read()
            except Exception as e:
                messagebox.showerror(APP_NAME, "读取失败：%s" % e)
                return
            ok, info, rs = LV.verify_license_bytes(raw)
            if not ok:
                messagebox.showerror("授权无效", rs)
                return
            import shutil
            dst = os.path.join(os.path.dirname(LV._license_path()), LV.LICENSE_FILENAME)
            try:
                shutil.copyfile(p, dst)
            except Exception as e:
                messagebox.showerror(APP_NAME, "无法写入授权文件：%s\n请手动把 license.key 放到程序所在目录。" % e)
                return
            messagebox.showinfo("授权成功", "授权已导入（授权给 %s，到期 %s）。\n重启软件后生效。"
                                % (info.get("name") or "（未填）", info.get("expiry") or "永久"))

        ttk.Button(f, text="导入授权文件 license.key…", command=import_license).grid(
            row=4, column=0, columnspan=2, sticky="w", pady=(2, 10))
        tip = ("使用步骤：\n"
               "1. 把上面的机器码复制后发给软件作者；\n"
               "2. 作者用“授权码生成器”签发与本机绑定的 license.key，发回给您；\n"
               "3. 点上面“导入授权文件”选择它（或把 license.key 放到本程序 exe 同一目录），重启软件即生效。\n\n"
               "说明：授权与本机机器码绑定，换机需重新申请；全过程离线，不联网。授权机制与防护边界见随附"
               "“软件授权说明.md”。")
        ttk.Label(f, text=tip, justify="left", foreground="#555", wraplength=760).grid(
            row=5, column=0, columnspan=3, sticky="w")

    def _tab_help(self):
        f = ttk.Frame(self.nb, padding=10)
        self.nb.add(f, text="说明")
        t = tk.Text(f, wrap="word", font=("Microsoft YaHei", 10))
        t.insert("1.0", HELP)
        t.configure(state="disabled")
        t.pack(fill="both", expand=True)

    # ------------------------------------------------------------------ 配置
    def get_cfg(self):
        cfg = copy.deepcopy(self.cfg)
        for k, v in self.iv.items():
            cfg[k] = v.get().strip()
        cfg["epsg"] = parse_epsg(cfg["epsg"])
        try:
            cfg["dpi"] = int(cfg["dpi"] or 300)
        except ValueError:
            cfg["dpi"] = 300
        cfg["formats"] = [k for k, v in self.fmt.items() if v.get()] or ["pdf"]
        for k, r in self.rows.items():
            cfg["layers"][k] = r.get()
        cfg["sensitive"] = [r.get() for r in self.sens_rows if r.v["path"].get().strip()]
        cfg["basemap"] = {k: v.get().strip() for k, v in self.bv.items()}
        f = cfg["fvc"]
        for k in ("ndvi", "red", "nir", "source"):
            f[k] = self.fv[k].get().strip()
        for k in ("scale", "offset", "s_pct", "v_pct"):
            f[k] = float(self.fv[k].get().strip() or C.FVC_DEFAULT[k])
        f["breaks"] = [float(x) for x in self.fv["breaks"].get().replace("，", ",").split(",") if x.strip()]
        f["labels"] = [x.strip() for x in self.fv["labels"].get().replace("，", ",").split(",") if x.strip()]
        if len(f["labels"]) != len(f["breaks"]) + 1:
            raise ValueError("植被覆盖度分级名称须比断点多 1 个（现有断点 %d 个、名称 %d 个）"
                             % (len(f["breaks"]), len(f["labels"])))
        if sorted(f["breaks"]) != f["breaks"] or any(not 0 < b < 1 for b in f["breaks"]):
            raise ValueError("植被覆盖度分级断点须在 0～1 之间并从小到大排列")
        if not 0 <= f["s_pct"] < f["v_pct"] <= 100:
            raise ValueError("NDVIs 百分位须小于 NDVIv 百分位，且在 0～100 之间")
        cfg["maps"] = [k for k, v in self.map_sel.items() if v]
        return cfg

    def set_cfg(self, cfg):
        cfg = C.merge_config(cfg)
        self.cfg = cfg
        for k, v in self.iv.items():
            v.set(str(cfg.get(k) if cfg.get(k) is not None else ""))
        e = cfg.get("epsg")
        self.iv["epsg"].set(next((s for s in EPSG_CHOICES if e and s.startswith(str(e))), str(e) if e else EPSG_CHOICES[0]))
        self.iv["scale_mode"].set(cfg.get("scale_mode") or "标准比例尺")
        for k, v in self.fmt.items():
            v.set(k in (cfg.get("formats") or []))
        for k, r in self.rows.items():
            r.clear()
            r.set(cfg["layers"].get(k, {}))
        for r in list(self.sens_rows):
            self.remove_sensitive(r)
        for s in cfg.get("sensitive") or []:
            self.add_sensitive(s)
        for k, v in self.bv.items():
            v.set(str((cfg.get("basemap") or {}).get(k) or ""))
        if not self.bv["tianditu_type"].get():
            self.bv["tianditu_type"].set("不使用")
        f = cfg["fvc"]
        for k in ("ndvi", "red", "nir", "scale", "offset", "s_pct", "v_pct", "source"):
            self.fv[k].set(str(f.get(k, "")))
        self.fv["breaks"].set(", ".join("%g" % b for b in f["breaks"]))
        self.fv["labels"].set(", ".join(f["labels"]))
        self.map_sel = {s["id"]: s["id"] in (cfg.get("maps") or []) for s in C.MAP_SPECS}
        self.refresh_maps()

    def new(self):
        if messagebox.askyesno(APP_NAME, "清空当前配置并新建？"):
            self.cfg_path = ""
            self.lbl_cfg.configure(text="未保存的配置")
            self.set_cfg(C.default_config())

    def open_cfg(self, path=None):
        p = path or filedialog.askopenfilename(filetypes=[("制图配置", "*.ecomap.json *.json")])
        if not p:
            return
        try:
            self.set_cfg(C.load_config(p))
            self.cfg_path = p
            self.lbl_cfg.configure(text=p)
            self.log("已打开配置：%s" % p)
        except Exception as e:
            messagebox.showerror(APP_NAME, "打开失败：%s" % e)

    def save_cfg(self):
        if not self.cfg_path:
            return self.save_cfg_as()
        try:
            C.save_config(self.get_cfg(), self.cfg_path)
            self.log("已保存配置：%s" % self.cfg_path)
        except Exception as e:
            messagebox.showerror(APP_NAME, str(e))

    def save_cfg_as(self):
        p = filedialog.asksaveasfilename(defaultextension=".ecomap.json", filetypes=[("制图配置", "*.ecomap.json")],
                                         initialfile="%s.ecomap.json" % (self.iv["project_name"].get() or "生态制图"))
        if p:
            self.cfg_path = p
            self.lbl_cfg.configure(text=p)
            self.save_cfg()

    def load_sample(self):
        import ecomap_sample
        d = os.path.join(tempfile.gettempdir(), "生态制图示例项目")
        try:
            p = ecomap_sample.make_sample(d)
        except Exception as e:
            messagebox.showerror(APP_NAME, "示例数据生成失败：%s" % e)
            return
        self.open_cfg(p)
        self.log("示例项目数据（虚构，仅供熟悉软件）位于：%s" % d)
        self.nb.select(4)

    def pick_out(self):
        p = filedialog.askdirectory()
        if p:
            self.iv["out_dir"].set(os.path.normpath(p))

    def _pick(self, var, types):
        p = filedialog.askopenfilename(filetypes=types)
        if p:
            var.set(os.path.normpath(p))

    # ------------------------------------------------------------------ 图件表
    def level_changed(self):
        lvl = self.iv["eval_level"].get() or "三级"
        self._select_maps("level")
        names = [C.SPEC_BY_ID[i]["title"] for i in C.required_maps(lvl)]
        self.log("评价等级改为%s：已勾选必须编制的 %d 幅图件（%s），其余图件请按项目情况自行勾选。"
                 % (lvl, len(names), "、".join(names)))

    def refresh_maps(self):
        lvl = self.iv["eval_level"].get() or "三级"
        tv = self.tv
        tv.delete(*tv.get_children())
        order = {"●": 0, "○": 1, "△": 2, "—": 3}
        specs = sorted(C.MAP_SPECS, key=lambda s: order.get(s["req"].get(lvl, "—"), 3))
        for s in specs:
            needs = []
            for n in s["needs"]:
                if n == "sensitive":
                    ok = bool(self.sens_rows and any(r.v["path"].get() for r in self.sens_rows))
                    needs.append(("生态敏感区", ok))
                elif n == "fvc":
                    ok = bool(self.fv["ndvi"].get() or (self.fv["red"].get() and self.fv["nir"].get()))
                    needs.append(("遥感影像", ok))
                else:
                    needs.append((C.SLOT_NAMES[n], bool(self.rows[n].v["path"].get().strip())))
            data = "、".join(("✓" if ok else "缺") + nm for nm, ok in needs)
            req = s["req"].get(lvl, "—")
            tag = "must" if req == "●" else "other"
            if tag == "must" and not all(ok for _, ok in needs):
                tag = "must_missing"
            tv.insert("", "end", iid=s["id"], tags=(tag,),
                      values=("☑" if self.map_sel.get(s["id"]) else "☐", s["title"],
                              "%s %s" % (req, C.REQ_TEXT[req]), s["basis"], data))
        tv.tag_configure("must", foreground="#000000", font=("Microsoft YaHei", 9, "bold"))
        tv.tag_configure("must_missing", foreground="#b00000", font=("Microsoft YaHei", 9, "bold"))
        tv.tag_configure("other", foreground="#777777")
        n = len(C.required_maps(lvl))
        self.lbl_req.configure(text="%s评价必须编制的图件 %d 幅（加粗，排在前面，缺资料的显示为红色），改变评价等级时自动勾选；"
                                    "其余图件（○涉及相应对象时编制、△可选、—导则未要求）请按项目情况自行勾选。"
                                    "分级系按 HJ 19—2022 正文条款归纳，地理位置图、水系图归为通用图件属推断，请复核。" % (lvl, n))

    def _toggle_map(self, e):
        if self.tv.identify_column(e.x) != "#1":
            return
        iid = self.tv.identify_row(e.y)
        if iid:
            self.map_sel[iid] = not self.map_sel.get(iid)
            self.tv.set(iid, "sel", "☑" if self.map_sel[iid] else "☐")

    def _select_maps(self, mode):
        lvl = self.iv["eval_level"].get() or "三级"
        req = set(C.required_maps(lvl))
        for s in C.MAP_SPECS:
            self.map_sel[s["id"]] = {"all": True, "none": False}.get(mode, s["id"] in req)
        self.refresh_maps()

    # ------------------------------------------------------------------ 配色
    def load_categories(self):
        key = self.color_keys[self.color_name.get()]
        try:
            cfg = self.get_cfg()
            g = C.read_layer(cfg["layers"][key], C.SLOT_NAMES[key])
        except Exception as e:
            messagebox.showerror(APP_NAME, str(e))
            return
        fld = cfg["layers"][key].get("field")
        if g is None or not fld:
            messagebox.showinfo(APP_NAME, "请先在“图层数据”页选择该专题的数据和分类字段")
            return
        import ecomap_render as R
        cats = sorted(g[fld].astype(str).unique())
        self.cfg = cfg
        cols = R.category_colors(cfg, key, cats)
        self._fill_colors(key, cols)

    def _fill_colors(self, key, cols):
        self.ctv.delete(*self.ctv.get_children())
        self._swatches = {}
        self._color_cur = (key, cols)
        for c, col in cols.items():
            img = tk.PhotoImage(width=40, height=14)
            img.put(col, to=(0, 0, 40, 14))
            self._swatches[c] = img
            self.ctv.insert("", "end", iid=c, image=img, values=(c, col))

    def edit_color(self, e):
        iid = self.ctv.identify_row(e.y)
        if not iid:
            return
        key, cols = self._color_cur
        rgb, hx = colorchooser.askcolor(color=cols[iid], title="类别“%s”的颜色" % iid)
        if hx:
            cols[iid] = hx
            self.cfg.setdefault("colors", {}).setdefault(key, {})[iid] = hx
            self._fill_colors(key, cols)

    def reset_colors(self):
        key = self.color_keys[self.color_name.get()]
        self.cfg.setdefault("colors", {}).pop(key, None)
        self.load_categories()

    # ------------------------------------------------------------------ 运行
    def log(self, msg):
        self.txt.insert("end", msg + "\n")
        self.txt.see("end")

    def _poll(self):
        try:
            while True:
                kind, *a = self.q.get_nowait()
                if kind == "log":
                    self.log(a[0])
                elif kind == "prog":
                    i, n, name = a
                    self.pb.configure(maximum=max(n, 1), value=i)
                    self.lbl_status.configure(text="%d/%d %s" % (i, n, name))
                elif kind == "done":
                    self.busy = False
                    self.lbl_status.configure(text=a[0])
                    if a[1]:
                        messagebox.showinfo(APP_NAME, a[1])
                elif kind == "error":
                    self.busy = False
                    self.lbl_status.configure(text="出错")
                    messagebox.showerror(APP_NAME, a[0])
                elif kind == "preview":
                    self._show_preview(*a)
        except queue.Empty:
            pass
        self.root.after(120, self._poll)

    def _start(self, fn):
        if self.busy:
            messagebox.showinfo(APP_NAME, "正在处理，请稍候")
            return
        try:
            cfg = self.get_cfg()
        except Exception as e:
            messagebox.showerror(APP_NAME, "参数有误：%s" % e)
            return
        self.cfg = cfg
        self.refresh_maps()
        self.busy = True
        self.nb.select(4)
        threading.Thread(target=fn, args=(cfg,), daemon=True).start()

    def _qlog(self, m):
        self.q.put(("log", m))

    def run_all(self):
        def work(cfg):
            import ecomap_run as RUN
            try:
                out, files, warns = RUN.run(cfg, log=self._qlog, progress=lambda i, n, s: self.q.put(("prog", i, n, s)))
                self.out_dir = out
                msg = "已生成 %d 个文件，输出目录：\n%s" % (len(files), out)
                if warns:
                    msg += "\n\n提示 %d 条（详见日志）：\n%s" % (len(warns), "\n".join(warns[:8]))
                self.q.put(("done", "完成", msg))
            except C.DataError as e:
                self.q.put(("error", str(e)))
            except Exception as e:
                self._qlog(traceback.format_exc())
                self.q.put(("error", "出错：%s" % e))
        self._start(work)

    def _get_ctx(self, cfg):
        key = json.dumps({k: v for k, v in cfg.items() if k not in ("maps", "colors", "out_dir", "formats", "dpi",
                                                                      "page", "scale_mode", "mapper", "project_name")},
                         ensure_ascii=False, sort_keys=True)
        if self.ctx_cache[0] == key:
            ctx = self.ctx_cache[1]
            ctx.cfg = C.merge_config(cfg)
            return ctx
        ctx = C.Context(cfg, log=self._qlog)
        self.ctx_cache = (key, ctx)
        return ctx

    def preview(self):
        sel = self.tv.selection()
        if not sel:
            messagebox.showinfo(APP_NAME, "请先在图件表中选中一行")
            return
        sid = sel[0]

        def work(cfg):
            import ecomap_render as R
            try:
                self.q.put(("prog", 0, 1, "读入数据"))
                ctx = self._get_ctx(cfg)
                if not all(ctx.has(n) for n in C.SPEC_BY_ID[sid]["needs"]):
                    self.q.put(("error", "该图件缺少数据：%s" % self.tv.set(sid, "data")))
                    return
                cv = R.render(ctx, sid, None)
                if cv is None:
                    self.q.put(("error", "该图件未能生成，详见日志"))
                    return
                for w in cv.warn:
                    self._qlog("提示：" + w)
                self.q.put(("preview", cv))
                self.q.put(("done", "预览完成", None))
            except C.DataError as e:
                self.q.put(("error", str(e)))
            except Exception as e:
                self._qlog(traceback.format_exc())
                self.q.put(("error", "出错：%s" % e))
        self._start(work)

    def _show_preview(self, cv):
        from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
        w = tk.Toplevel(self.root)
        w.title("预览：%s（比例尺 1:%d）" % (cv.title, cv.scale))
        w.geometry("820x1000" if cv.L["H"] > cv.L["W"] else "1150x860")
        cv.fig.set_dpi(80)
        canvas = FigureCanvasTkAgg(cv.fig, master=w)
        NavigationToolbar2Tk(canvas, w)
        canvas.get_tk_widget().pack(fill="both", expand=True)
        canvas.draw()

    def open_out(self):
        d = getattr(self, "out_dir", None) or self.iv["out_dir"].get()
        if d and os.path.isdir(d):
            if sys.platform == "win32":
                os.startfile(d)
        else:
            messagebox.showinfo(APP_NAME, "输出目录尚不存在")


# ====================================================================== 自检
def selftest(path):
    import time
    lines, ok = [], True

    def w(s):
        lines.append(s)

    t0 = time.time()
    w("%s 自检  %s" % (TITLE, time.strftime("%Y-%m-%d %H:%M:%S")))
    try:
        import ecomap_render as R
        import ecomap_run as RUN
        import ecomap_sample
        import geopandas, pyogrio, rasterio, shapely, pyproj, matplotlib
        w("库版本：geopandas %s，shapely %s，pyproj %s（PROJ %s），pyogrio %s（GDAL %s），rasterio %s（GDAL %s），matplotlib %s"
          % (geopandas.__version__, shapely.__version__, pyproj.__version__, pyproj.proj_version_str,
             pyogrio.__version__, pyogrio.__gdal_version_string__, rasterio.__version__, rasterio.__gdal_version__,
             matplotlib.__version__))
        w("中文字体：%s" % ("、".join(R.FONTS) if R.FONTS else "未找到（图面中文将显示为方框）"))
        # 纯计算
        assert C.cgcs2000_3deg_epsg(118.6) == 4549 and C.cgcs2000_3deg_epsg(117.9) == 4548
        assert C.choose_scale(3000, 3000, 180, 200) == 25000
        lon, lat = C.gcj02_to_wgs84(*C.wgs84_to_gcj02(118.55, 25.05))
        assert abs(lon - 118.55) < 1e-7 and abs(lat - 25.05) < 1e-7
        w("纯计算校验：通过")
        d = os.path.join(tempfile.mkdtemp(prefix="ecomap_selftest_"), "示例")
        cfgp = ecomap_sample.make_sample(d)
        cfg = C.load_config(cfgp)
        out, files, warns = RUN.run(cfg, log=lambda m: w("  " + m))
        bad = [p for _, p in files if not os.path.exists(p) or os.path.getsize(p) < 5000]
        n_expect = 2 * 18
        w("输出文件 %d 个（应为 %d 个），异常 %d 个" % (len(files), n_expect, len(bad)))
        if len(files) != n_expect or bad:
            ok = False
        xl = os.path.join(out, "生态图件清单及面积统计.xlsx")
        if not os.path.exists(xl):
            ok = False
            w("缺少 %s" % xl)
        from openpyxl import load_workbook
        wb = load_workbook(xl)
        w("Excel 工作表：%s" % "、".join(wb.sheetnames))
        if any("生成失败" in x for x in warns):
            ok = False
        # 界面
        root = tk.Tk()
        app = App(root)
        app.open_cfg(cfgp)
        app.refresh_maps()
        c2 = app.get_cfg()
        assert c2["layers"]["landuse"]["field"] == "DLMC"
        root.update()
        root.destroy()
        w("界面构建与配置读写：通过")
        # 授权机制自检：机器码稳定、演示私钥签发的授权可被对应公钥验证、换机/篡改被拒
        import importlib
        import license_verify as _LV
        sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "tools"))
        code = _LV.machine_code()
        assert code == _LV.machine_code() and "-" in code
        try:
            import license_issuer as _LI
            from cryptography.hazmat.primitives import serialization as _ser
            demo_key = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tools", "keys_demo",
                                    "private_demo.pem")
            with open(demo_key, "rb") as fh:
                dpub = _ser.load_pem_private_key(fh.read(), password=None).public_key().public_bytes(
                    _ser.Encoding.Raw, _ser.PublicFormat.Raw).hex()
            text, _ = _LI.make_license(code, "自检", None, demo_key)
            okL, _, _ = _LV.verify_license_bytes(text, code, public_key_hex=dpub)
            okM, _, rM = _LV.verify_license_bytes(text, "AAAA-BBBB-CCCC-DDDD", public_key_hex=dpub)
            w("授权机制：机器码 %s；签发/验证 %s；换机拒绝 %s" % (code, "通过" if okL else "失败",
                                                      "通过" if (not okM) else "失败"))
            if not (okL and not okM):
                ok = False
        except ImportError:
            w("授权机制：机器码 %s（打包版不含签发工具 cryptography，跳过签发校验）" % code)
        w("示例输出目录：%s" % out)
        dst = os.environ.get("ECOMAP_SAMPLE_OUT")
        if dst:   # 构建流程用：把示例成果复制到指定目录
            import shutil
            shutil.copytree(out, dst, dirs_exist_ok=True)
            for suffix, alias in (("_土地利用现状图.jpg", "sample_landuse.jpg"),
                                  ("_工程占用土地利用类型叠置图.jpg", "sample_occupy.jpg")):
                fn = next(f for f in os.listdir(out) if f.endswith(suffix))
                shutil.copy(os.path.join(out, fn), os.path.join(dst, alias))
    except Exception:
        ok = False
        w(traceback.format_exc())
    w("用时 %.1f s；自检%s" % (time.time() - t0, "通过" if ok else "失败"))
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return 0 if ok else 1


def main():
    if len(sys.argv) >= 3 and sys.argv[1] == "--selftest":
        sys.exit(selftest(sys.argv[2]))
    if "--print-machine" in sys.argv:
        print(LV.machine_code())
        sys.exit(0)
    if len(sys.argv) >= 3 and sys.argv[1] == "--batch":
        # 命令行批量出图同样受授权约束
        state, _, reason = LV.license_state()
        if state == "expired":
            print("未授权，试用已用完：%s（机器码 %s）" % (reason, LV.machine_code()))
            sys.exit(2)
        import ecomap_run as RUN
        RUN.run(C.load_config(sys.argv[2]))
        return
    if sys.platform == "win32":
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    # 启动门禁：已授权直接运行；未授权则给予首次试用，到时弹授权窗并锁定；试用用完则直接要求授权
    demo_mode = any(a.startswith("--demo") for a in sys.argv)   # 截图/演示模式：静默试用，不弹试用提示框
    state, payload, reason = LV.license_state()
    if state == "expired" and not demo_mode:
        LV.show_gate(None, "试用时间已用完（共 %d 分钟），请授权后继续使用。" % LV.TRIAL_MINUTES)
        return
    root = tk.Tk()
    app = App(root)
    if "--demo-license" in sys.argv:
        app.nb.select(7)
    elif state == "trial" and not demo_mode:
        remaining = max(1, int(payload))
        mins = max(1, (remaining + 59) // 60)
        try:
            messagebox.showinfo("试用", "未检测到授权，您可先试用本软件。\n"
                                "本次试用剩余约 %d 分钟，到时将提示授权。\n"
                                "如需长期使用，请在“软件授权”页复制机器码发给作者申请。" % mins)
        except Exception:
            pass

        def _on_trial_end():
            try:
                LV.show_gate(root, "试用时间已到（共 %d 分钟），请授权后继续使用。" % LV.TRIAL_MINUTES)
            finally:
                try:
                    root.destroy()
                except Exception:
                    pass

        root.after(remaining * 1000, _on_trial_end)
    if "--demo-maps" in sys.argv:
        app.iv["eval_level"].set("三级")
        app.level_changed()
        app.nb.select(4)
    elif "--demo-data" in sys.argv:
        app.nb.select(6)
    elif "--demo-layers" in sys.argv:
        app.load_sample()
        app.nb.select(1)
    elif "--demo" in sys.argv:
        app.load_sample()
        sid = "landuse"
        app.tv.selection_set(sid)
        root.after(1500, app.preview)
    elif len(sys.argv) >= 2 and os.path.exists(sys.argv[1]):
        app.open_cfg(sys.argv[1])
    root.mainloop()


if __name__ == "__main__":
    main()
