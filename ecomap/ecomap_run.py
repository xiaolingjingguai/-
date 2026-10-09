# -*- coding: utf-8 -*-
"""批量出图与成果导出（图件、面积统计 Excel、图件清单、制图日志）。"""
import datetime
import os
import traceback

import ecomap_core as C
import ecomap_render as R


def write_excel(path, ctx, check_rows, files):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    wb = Workbook()
    thin = Side(style="thin")
    head_font = Font(name="宋体", bold=True, size=11)
    body_font = Font(name="宋体", size=10.5)

    def sheet(ws, title, header, rows, widths):
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
        ws["A1"].font = Font(name="黑体", size=12)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.append(header)
        for r in rows:
            ws.append(list(r))
        n = ws.max_row
        for row in ws.iter_rows(min_row=2, max_row=n):
            for c in row:
                c.font = head_font if c.row == 2 else body_font
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                top = Side(style="medium") if c.row == 2 else None
                bottom = Side(style="medium") if c.row in (2, n) else None
                c.border = Border(top=top, bottom=bottom)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w

    ws = wb.active
    ws.title = "图件清单"
    sheet(ws, "生态影响评价图件清单（评价等级：%s）" % ctx.cfg["eval_level"],
          ["序号", "图件名称", "要求", "HJ 19—2022 条款", "本次情况"], check_rows, [6, 34, 6, 34, 30])
    ws.append([])
    ws.append(["注：●应编制；○涉及相应对象时编制；△可选；—导则未要求。分级为按导则正文条款归纳，序号1、2归为通用图件属推断，请复核。"])
    for name, rows in ctx.stats.items():
        ws = wb.create_sheet(name[:31])
        sheet(ws, name, ["类型", "面积（hm²）", "比例（%）"], rows, [30, 16, 12])
        ws.append([])
        if name.startswith("工程占用"):
            ws.append(["注：图形叠置法（HJ 19—2022 8.2.1、附录C.2），比例为占工程占地内该专题数据覆盖面积的百分比。"])
        elif "植被覆盖度" in name:
            f = ctx.fvc
            ws.append(["注：式C.5，NDVIs=%.4f（%s%%分位）、NDVIv=%.4f（%s%%分位），分级断点 %s；百分位与断点为经验取值，导则未规定。"
                       % (f["ns"], ctx.cfg["fvc"]["s_pct"], f["nv"], ctx.cfg["fvc"]["v_pct"],
                          "、".join("%g%%" % (b * 100) for b in ctx.cfg["fvc"]["breaks"]))])
        else:
            ws.append(["注：面积按 %s 平面坐标计算。" % C.crs_desc(ctx.crs)])
    ws = wb.create_sheet("输出文件")
    sheet(ws, "输出文件", ["图件", "文件"], files, [40, 70])
    wb.save(path)


def run(cfg, ids=None, log=print, progress=None):
    """按配置生成全部图件。返回 (输出目录, 生成文件列表, 警告列表)。"""
    out_dir = cfg.get("out_dir") or os.path.join(os.path.expanduser("~"), "生态图件输出")
    os.makedirs(out_dir, exist_ok=True)
    lines = []

    def L(msg):
        lines.append(msg)
        log(msg)

    L("%s 生态制图 %s 开始：%s" % (C.STD, C.VERSION, datetime.datetime.now().strftime("%Y-%m-%d %H:%M")))
    if not R.FONTS:
        L("警告：未找到中文字体（黑体/微软雅黑/宋体），图面中文可能显示为方框")
    ctx = C.Context(cfg, log=L)
    todo = R.jobs(ctx, ids)
    files, warns, generated = [], [], {}
    for k, (sid, si, fname) in enumerate(todo, 1):
        if progress:
            progress(k - 1, len(todo), fname)
        try:
            cv = R.render(ctx, sid, si)
            if cv is None:
                continue
            ps = cv.save(os.path.join(out_dir, fname), cfg.get("formats") or ["pdf"], int(cfg.get("dpi") or 300))
            generated.setdefault(sid, []).append(fname)
            for p in ps:
                files.append((fname, p))
            L("已生成：%s（比例尺 1:%d）" % (fname, cv.scale))
            for w in cv.warn:
                warns.append("%s：%s" % (fname, w))
                L("  提示：" + w)
        except Exception as e:
            warns.append("%s：生成失败 %s" % (fname, e))
            L("生成失败：%s\n%s" % (fname, traceback.format_exc()))
    if progress:
        progress(len(todo), len(todo), "完成")
    rows = C.checklist(ctx, generated)
    for r in rows:
        if r[2] in ("●",) and not r[4].startswith("已生成") and "本程序不生成" not in r[4]:
            L("注意：%s 为%s应编制图件，%s" % (r[1], ctx.cfg["eval_level"], r[4]))
    xlsx = os.path.join(out_dir, "生态图件清单及面积统计.xlsx")
    try:
        write_excel(xlsx, ctx, rows, files)
        L("已导出：%s" % xlsx)
    except PermissionError:
        L("导出失败：%s 被占用（请关闭 Excel 后重试）" % xlsx)
    C.save_config(ctx.cfg, os.path.join(out_dir, "制图配置.ecomap.json"))
    with open(os.path.join(out_dir, "制图日志.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return out_dir, files, warns
