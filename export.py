# -*- coding: utf-8 -*-
"""导出模块：Excel（汇总表 + 逐样明细）与 Word（三线表）"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

import re
from gb3838_limits import CLASS_CN, STANDARD_INFO

_SUP = str.maketrans("0123456789-+", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺")


def sup_unicode(v):
    """Excel 用：4.7×10^3 → 4.7×10³"""
    if not isinstance(v, str):
        return v
    return re.sub(r"\^([-+]?\d+)", lambda m: m.group(1).translate(_SUP), v)


def add_runs(para, text, size, bold, red=False):
    """Word 用：把 ^-5 写成上标 run；red=True 时为红字"""
    for k, part in enumerate(re.split(r"\^([-+]?\d+)", text)):
        if part:
            r = para.add_run(part)
            _font(r, size, bold)
            if red:
                r.font.color.rgb = RGBColor(0xFF, 0, 0)
            if k % 2 == 1:
                r.font.superscript = True

SUMMARY_HEAD = ["监测断面", "监测项目", "单位", "样次", "监测值范围", "标准限值",
                "最大标准指数", "超标率（%）", "最大超标倍数", "单因子水质类别", "达标情况", "备注"]
SUMMARY_KEYS = ["station", "item", "unit", "n", "range", "limit",
                "max_index", "exceed_rate", "max_exceed", "cls", "result", "note"]
DETAIL_HEAD = ["监测断面", "采样日期", "监测项目（报告原名）", "对应标准项目", "所在表", "监测值", "单位",
               "标准限值", "标准指数", "超标倍数", "单因子水质类别", "判定结果", "备注"]
DETAIL_KEYS = ["station", "date", "item", "std_name", "table", "value", "unit",
               "limit_text", "index", "exceed_times", "cls", "result", "note"]


def table_title(target_class, water_body):
    wb = "湖、库" if water_body == "lake" else "河流"
    return f"地表水环境现状监测结果及评价（{wb}，执行 {STANDARD_INFO['code']} {CLASS_CN[target_class]}标准）"


def table_note(target_class):
    return ("注：1. 评价标准为《地表水环境质量标准》（GB 3838-2002）"
            f"{CLASS_CN[target_class]}标准；\n"
            "2. 标准指数 S=C/Cs，pH、溶解氧按 HJ 2.3-2018 水质指数法计算，S＞1 为超标；\n"
            "3. 超标倍数=(C-Cs)/Cs，水温、pH、溶解氧不计算超标倍数；\n"
            "4. 未检出项目以\"检出限+L\"表示，检出限不高于标准限值的判定为达标；\n"
            "5. 超标率=超标样次/有效判定样次×100%。")


# ------------------------------------------------------------------ 超标标红
RED_HEX = "FF0000"
_EXCEED_KEYS = {"value", "value_disp", "range", "index", "max_index", "exceed_rate", "max_exceed", "exceed_times",
                "cls", "result", "ratio", "max_ratio", "over"}


def red_keys(row):
    """该行需要标红（红字）的字段：超标行的监测值、指数、超标率/倍数、类别、结论"""
    if "_red" in row:
        return set(row["_red"])
    if str(row.get("result", "")).startswith("超"):      # 超标 / 超筛选值 / 超管制值
        return _EXCEED_KEYS
    if "det_rate" in row:                       # 地下水统计表：超标率＞0 标红
        try:
            if float(str(row.get("exceed_rate", "0")).strip("%")) > 0:
                return {"exceed_rate", "max"}
        except ValueError:
            pass
    return set()


# ------------------------------------------------------------------ 通用写出
def write_excel(path, sheets):
    """sheets: list of (工作表名, 表题, 表头, 键, 数据行, 表注)"""
    wb = Workbook()
    thick, thin = Side(style="medium"), Side(style="thin")
    for k_sheet, (name, title, head, keys, data, note_text) in enumerate(sheets):
        ws = wb.active if k_sheet == 0 else wb.create_sheet()
        ws.title = name
        ncol = len(head)
        ws.cell(1, 1, title)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws.cell(1, 1).font = Font(name="宋体", size=11, bold=True)
        ws.cell(1, 1).alignment = Alignment(horizontal="center")
        for j, h in enumerate(head, 1):
            ws.cell(2, j, h).font = Font(name="宋体", size=10.5, bold=True)
        for i, r in enumerate(data, 3):
            for j, k in enumerate(keys, 1):
                ws.cell(i, j, sup_unicode(r.get(k, "")))
        last = 2 + len(data)
        for i in range(2, last + 1):
            reds = red_keys(data[i - 3]) if i > 2 else set()
            for j in range(1, ncol + 1):
                c = ws.cell(i, j)
                if i > 2:
                    c.font = Font(name="宋体", size=10.5, color=RED_HEX if keys[j - 1] in reds else None)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c.border = Border(top=thick if i == 2 else None,
                                  bottom=thick if i == last else (thin if i == 2 else None))
        if note_text:
            note = ws.cell(last + 1, 1, note_text)
            ws.merge_cells(start_row=last + 1, start_column=1, end_row=last + 1, end_column=ncol)
            note.alignment = Alignment(wrap_text=True, vertical="top")
            note.font = Font(name="宋体", size=9)
            ws.row_dimensions[last + 1].height = 15 * (note_text.count("\n") + 2)
        for j, h in enumerate(head, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(10, len(h) * 2 + 2)
        ws.column_dimensions[get_column_letter(ncol)].width = max(
            ws.column_dimensions[get_column_letter(ncol)].width, 24)
    wb.save(path)


# ------------------------------------------------------------------ Word
def _set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, sz in kw.items():
        el = OxmlElement(f"w:{edge}")
        if sz:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)


def _font(run, size=10.5, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _word_table(doc, title, head, keys, data, note_text, merge_first=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("表  " + title), 10.5, True)
    n = len(data)
    t = doc.add_table(rows=1 + n, cols=len(keys))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 三线表：表级只保留上、下框线（1.5 磅），表头下框线 0.75 磅
    tb = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in ("top", "bottom"):
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "12"); el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        tb.append(el)
    t._tbl.tblPr.append(tb)
    for i in range(n + 1):
        reds = red_keys(data[i - 1]) if i else set()
        for j, k in enumerate(keys):
            cell = t.cell(i, j)
            text = head[j] if i == 0 else str(data[i - 1].get(k, ""))
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(para, text, 9, i == 0, k in reds)
            if i == 0:
                _set_cell_border(cell, bottom=6)
    if merge_first and n:
        k0 = keys[0]
        start = 1
        for i in range(2, n + 2):
            if i == n + 1 or data[i - 1][k0] != data[start - 1][k0]:
                if i - 1 > start:
                    merged = t.cell(start, 0).merge(t.cell(i - 1, 0))
                    for extra in merged.paragraphs[1:]:
                        extra._element.getparent().remove(extra._element)
                start = i
    for line in (note_text or "").split("\n"):
        if line:
            np_ = doc.add_paragraph()
            np_.paragraph_format.space_after = Pt(0)
            _font(np_.add_run(line), 9)
    doc.add_paragraph()


def write_word(path, tables):
    """tables: list of (表题, 表头, 键, 数据行, 表注)"""
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    for title, head, keys, data, note_text in tables:
        _word_table(doc, title, head, keys, data, note_text)
    doc.save(path)


def collect_notes(rows, start_no):
    notes = list(dict.fromkeys(n for r in rows for n in r.get("note", "").split("；") if n))
    return f"\n{start_no}. 其他说明：" + "；".join(notes) + "。" if notes else ""


# ------------------------------------------------------------------ 地表水
def to_excel(path, summary, detail, target_class, water_body):
    title = table_title(target_class, water_body)
    note = table_note(target_class)
    write_excel(path, [("汇总表", title, SUMMARY_HEAD, SUMMARY_KEYS, summary, note),
                       ("逐样明细", title, DETAIL_HEAD, DETAIL_KEYS, detail, note)])


def to_word(path, summary, target_class, water_body, cols=None):
    """cols：导出到 Word 的列（默认去掉"备注"，备注并入表注）"""
    keys = cols or [k for k in SUMMARY_KEYS if k != "note"]
    head = [SUMMARY_HEAD[SUMMARY_KEYS.index(k)] for k in keys]
    note = table_note(target_class) + collect_notes(summary, 6)
    write_word(path, [(table_title(target_class, water_body), head, keys, summary, note)])


# ------------------------------------------------------------------ 地下水
GW_SUMMARY_HEAD = ["监测点位", "监测指标", "单位", "样次", "监测值范围", "标准限值",
                   "最大标准指数", "超标率（%）", "最大超标倍数", "单指标类别", "达标情况", "备注"]
GW_DETAIL_HEAD = ["监测点位", "采样日期", "监测指标（报告原名）", "对应标准指标", "所在表", "监测值", "单位",
                  "标准限值", "标准指数", "超标倍数", "单指标类别", "判定结果", "备注"]
GW_OVERALL_HEAD = ["监测点位", "采样日期", "参评指标数", "综合类别", "最差类别指标", "未检出且类别未定指标"]
GW_OVERALL_KEYS = ["station", "date", "n", "cls", "worst_items", "nd_items"]


def gw_title(target_class):
    return f"地下水环境现状监测结果及评价（执行 GB/T 14848-2017 {CLASS_CN[target_class]}标准）"


def gw_note(target_class):
    return ("注：1. 评价标准为《地下水质量标准》（GB/T 14848-2017）"
            f"{CLASS_CN[target_class]}标准；\n"
            "2. 单指标类别按 GB/T 14848-2017 第6.2条确定，限值相同时从优不从劣；\n"
            "3. 标准指数按 HJ 610-2016 第8.4.1.2条公式（2）~（4）计算，P＞1 为超标；超标倍数=(C-Cs)/Cs；\n"
            "4. 未检出项目以\"检出限+L\"表示；检出限高于Ⅰ类限值的记\"≤X类\"，表示只能证明不劣于该类别；\n"
            "5. 超标率=超标样次/有效判定样次×100%。")


def gw_overall_note():
    return ("注：综合类别按 GB/T 14848-2017 第6.3条，取单指标评价结果最差的类别；"
            "未检出且检出限高于Ⅰ类限值的指标（记\"≤X类\"）不参与综合类别判定，单独列出。")


GW_STAT_HEAD = ["监测指标", "单位", "样本数", "检出数", "检出率（%）", "最小值", "最大值", "均值", "标准差", "超标率（%）"]
GW_STAT_KEYS = ["item", "unit", "n", "n_det", "det_rate", "min", "max", "mean", "sd", "exceed_rate"]


def gw_stat_note():
    return ("注：按 HJ 610-2016 第8.4.1.1条进行统计分析；最小值、最大值、均值、标准差仅以检出值计算，"
            "标准差为样本标准差，检出值少于2个时记\"—\"。")


def _mark_overall(rows, target_class):
    from gw_judge import RANK
    lim = RANK.get(CLASS_CN[target_class])
    out = []
    for r in rows:
        r = dict(r)
        rk = RANK.get(r.get("cls"))
        r["_red"] = {"cls", "worst_items"} if (rk is not None and lim is not None and rk > lim) else set()
        out.append(r)
    return out


def gw_to_excel(path, summary, detail, overall_rows, target_class, stat_rows=()):
    overall_rows = _mark_overall(overall_rows, target_class)
    title = gw_title(target_class)
    note = gw_note(target_class)
    write_excel(path, [
        ("汇总表", title, GW_SUMMARY_HEAD, SUMMARY_KEYS, summary, note),
        ("综合评价", "地下水质量综合评价结果", GW_OVERALL_HEAD, GW_OVERALL_KEYS, overall_rows, gw_overall_note()),
        ("统计分析", "地下水水质现状监测结果统计", GW_STAT_HEAD, GW_STAT_KEYS, list(stat_rows), gw_stat_note()),
        ("逐样明细", title, GW_DETAIL_HEAD, DETAIL_KEYS, detail, note)])


def gw_to_word(path, summary, overall_rows, target_class, stat_rows=()):
    overall_rows = _mark_overall(overall_rows, target_class)
    keys = [k for k in SUMMARY_KEYS if k != "note"]
    head = [GW_SUMMARY_HEAD[SUMMARY_KEYS.index(k)] for k in keys]
    note = gw_note(target_class) + collect_notes(summary, 6)
    write_word(path, [
        (gw_title(target_class), head, keys, summary, note),
        ("地下水水质现状监测结果统计", GW_STAT_HEAD, GW_STAT_KEYS, list(stat_rows), gw_stat_note()),
        ("地下水质量综合评价结果", GW_OVERALL_HEAD, GW_OVERALL_KEYS, overall_rows, gw_overall_note())])
