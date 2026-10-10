# -*- coding: utf-8 -*-
"""
按用户提供的第三方报告截图（2026.07.22，两处山涧水断面）手工转录，生成模拟"原始 Word 报告"，
用于调试 Word 解析。指数上标在截图中模糊，转录值（10³、10⁻⁵、10⁻⁴、10⁻³、10⁻³）为人工辨读，须以原件为准。
"""
import os
from docx import Document

ROWS = [  # 项目, 1#, 2#, 标准值；值中 ^x 表示上标
    ("pH（无量纲）", "7.2(26℃)", "7.2(28℃)", "6-9"),
    ("氨氮（mg/L）", "0.70", "0.32", "≤1.0"),
    ("总磷（mg/L）", "0.19", "0.26", "≤0.2"),
    ("总氮（mg/L）", "3.76", "0.92", "≤1.0"),
    ("COD（mg/L）", "17", "11", "≤20"),
    ("BOD5（mg/L）", "3.8", "2.2", "≤4.0"),
    ("粪大肠杆菌（MPN/L）", "4.7×10^3", "3.3×10^3", "≤10000"),
    ("铜（mg/L）", "0.05L", "0.05L", "≤1.0"),
    ("锌（mg/L）", "0.09", "0.05L", "≤1.0"),
    ("汞（mg/L）", "4.00×10^-5L", "4.00×10^-5L", "≤0.0001"),
    ("砷（mg/L）", "3.0×10^-4L", "3.0×10^-4L", "≤0.05"),
    ("铅（mg/L）", "1.0×10^-3L", "1.0×10^-3L", "≤0.05"),
    ("镉（mg/L）", "0.1×10^-3L", "0.1×10^-3L", "≤0.005"),
    ("高锰酸盐指数（mg/L）", "5.6", "3.9", "≤6"),
]


def put(cell, text):
    p = cell.paragraphs[0]
    if "^" in text:
        base, rest = text.split("^", 1)
        exp = rest.rstrip("L")
        p.add_run(base)
        p.add_run(exp).font.superscript = True
        if rest.endswith("L"):
            p.add_run("L")
    else:
        p.add_run(text)


doc = Document()
t = doc.add_table(rows=3 + len(ROWS) + 1, cols=5)
t.cell(0, 0).merge(t.cell(2, 0)); put(t.cell(0, 0), "检测日期")
t.cell(0, 1).merge(t.cell(2, 1)); put(t.cell(0, 1), "检测项目")
t.cell(0, 2).merge(t.cell(0, 3)); put(t.cell(0, 2), "监测点位、样品编号及结果")
t.cell(0, 4).merge(t.cell(2, 4)); put(t.cell(0, 4), "标准值或范围")
put(t.cell(1, 2), "蛋鸡养殖场周边山涧水系1#"); put(t.cell(1, 3), "育雏养殖场周边山涧水系2#")
put(t.cell(2, 2), "S2607222-1-1"); put(t.cell(2, 3), "S2607222-2-1")
t.cell(3, 0).merge(t.cell(2 + len(ROWS), 0)); put(t.cell(3, 0), "2026.07.22")
for i, r in enumerate(ROWS, 3):
    for j, v in enumerate(r, 1):
        put(t.cell(i, j), v)
last = 3 + len(ROWS)
t.cell(last, 0).merge(t.cell(last, 4))
put(t.cell(last, 0), "备注：执行《地表水环境质量标准》（GB3838-2002）III类水质标准。")
doc.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "样例_山涧水_模拟原件.docx"))
